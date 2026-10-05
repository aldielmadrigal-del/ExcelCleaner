from openpyxl import Workbook
from datetime import datetime

wb = Workbook()

# =========================================================
# HOJA 1 — Clientes: muchos problemas mezclados
# =========================================================

ws = wb.active
ws.title = "Clientes"

ws.append(["Nombre", "Edad", "Ciudad", "Email", "Telefono", "Fecha"])
ws.append([" Juan ", 20, " Camaguey ", " juan@gmail.com ", " 555-1000 ", datetime(2026, 1, 15)])
ws.append(["Maria", 25, "Florida", "maria@gmail.com", "555-2000", datetime(2026, 2, 10)])
ws.append(["Juan ", 20, "Camaguey", "juan@gmail.com", "555-1000", datetime(2026, 1, 15)])
ws.append([])
ws.append([" Pedro  ", 31, " Holguin", "pedro@gmail.com", "555-3000", datetime(2026, 3, 5)])
ws.append(["Pedro", 31, "Holguin", "pedro@gmail.com", "555-3000", datetime(2026, 3, 5)])
ws.append(["Ana", 22, "Santiago ", " ana@gmail.com", "555-4000 ", datetime(2026, 4, 20)])
ws.append(["Ana", 22, "Santiago ", " ana@gmail.com", "555-4000", datetime(2026, 4, 20)])
ws.append(["Luis", None, " Havana ", "luis@gmail.com ", "555-5000", None])
ws.append(["  Carlos  ", 40, "Cienfuegos", "carlos@gmail.com", "555-6000", datetime(2026, 5, 1)])
ws.append(["Carlos", 40, "Cienfuegos", "carlos@gmail.com", "555-6000", datetime(2026, 5, 1)])
ws.append(["Texto raro", "20 años", "  Camaguey  ", " texto ", " 123 ", "fecha desconocida"])
ws.append(["Con espacios internos", 33, "Santa Clara", "ana @gmail.com", "555 7000", datetime(2026, 6, 1)])

# Otra fila completamente vacía
ws.append([])

# =========================================================
# HOJA 2 — Ventas
# =========================================================

ws = wb.create_sheet("Ventas")

ws.append(["Producto", "Cantidad", "Precio", "Total", "Vendedor"])
ws.append([" Laptop ", 2, 720, 1440, " Juan "])
ws.append(["Laptop", 2, 720, 1440, "Juan"])
ws.append([" Mouse", 5, 20, 100, "Maria"])
ws.append(["Mouse ", 5, 20, 100, "Maria "])
ws.append([])
ws.append(["Teclado", 10, 15, 150, "Pedro"])
ws.append(["Teclado", 10, 15, 150, "Pedro"])
ws.append(["Monitor", None, 200, None, " Ana "])
ws.append(["Monitor", None, 200, None, " Ana "])
ws.append(["Producto desconocido", "10 unidades", "20 USD", "200 USD", "Carlos"])

# =========================================================
# HOJA 3 — Inventario
# =========================================================

ws = wb.create_sheet("Inventario")

ws.append(["Producto", "Stock", "Categoria"])
ws.append([" Teclado ", 15, "Perifericos"])
ws.append(["Teclado", 15, "Perifericos"])
ws.append(["Monitor", 8, " Pantallas "])
ws.append(["Monitor ", 8, "Pantallas"])
ws.append(["Laptop", 5, "Computadoras"])
ws.append(["Laptop", 5, "Computadoras"])
ws.append([])
ws.append(["USB", 0, "Almacenamiento"])
ws.append(["USB", 0, "Almacenamiento"])

# =========================================================
# HOJA 4 — Sin errores
# =========================================================

ws = wb.create_sheet("SinErrores")

ws.append(["ID", "Nombre", "Edad"])
ws.append([1, "Carlos", 30])
ws.append([2, "Maria", 25])
ws.append([3, "Pedro", 40])
ws.append([4, "Ana", 22])

# =========================================================
# HOJA 5 — Casi vacía
# =========================================================

ws = wb.create_sheet("Vacia")

ws.append([])
ws.append([])
ws.append([])

# =========================================================
# HOJA 6 — Casos difíciles
# =========================================================

ws = wb.create_sheet("CasosDificiles")

ws.append(["Dato", "Valor"])
ws.append(["Cero", 0])
ws.append(["Negativo", -50])
ws.append(["Decimal", 12.50])
ws.append(["Texto numerico", "00125"])
ws.append(["Codigo", " 000123 "])
ws.append(["Mayusculas", "JUAN"])
ws.append(["Minusculas", "juan"])
ws.append(["Espacios internos", "Juan Carlos"])
ws.append(["Muchos espacios internos", "Juan   Carlos"])
ws.append(["Email raro", "usuario+test@gmail.com"])
ws.append(["URL", " https://example.com "])
ws.append(["Texto vacio", ""])
ws.append(["NULL como texto", "NULL"])
ws.append(["NA como texto", "N/A"])
ws.append(["Guion", "-"])

# Guardar
wb.save("examples/test_chaos.xlsx")

print("Archivo creado: examples/test_chaos.xlsx")
print("Hojas:", wb.sheetnames)
